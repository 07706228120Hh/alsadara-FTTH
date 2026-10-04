# -*- coding: utf-8 -*-
import csv, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

base = r"C:\SadaraPlatform\docs\reports"
out = os.path.join(base, "مشاكل_محاسبية_Sadara_2026-08-29.xlsx")

sheets = [
    ("قيود يدوية غير متوازنة", "ds1_manual.csv",
     "أخطاء إدخال يدوية (مايو 2026) — قيود تسوية حسابات أُدخِلت بمدين≠دائن"),
    ("تعديل اشتراك غير متوازن", "ds2_subscription.csv",
     "قيود تجديد/شراء اشتراك على فني عُدِّلت فاختلّ توازنها (بق في مسار التعديل)"),
    ("أجور مهام غير مُرحّلة", "ds3_unposted_fees.csv",
     "أجور 1000 د.ع سُجِّلت على الفني بلا قيد محاسبي (ضحايا نافذة تيتيم CompanyId)"),
]

hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(color="FFFFFF", bold=True, size=11)
title_font = Font(bold=True, size=13, color="1F3864")
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
num_cols = {"مدين", "دائن", "الفرق", "المبلغ"}

wb = openpyxl.Workbook()

# ---- ورقة الملخص ----
ws = wb.active
ws.title = "ملخص"
ws.sheet_view.rightToLeft = True
ws["A1"] = "مشاكل محاسبية — Sadara Company"
ws["A1"].font = Font(bold=True, size=15, color="1F3864")
ws["A2"] = "أُنشئ: 2026-08-29 — الدفتر غير متوازن بـ 2,165,000 د.ع"
ws["A2"].font = Font(size=11, color="C00000", bold=True)
summ = [
    ["المجموعة", "عدد", "الفرق (د.ع)", "التشخيص"],
    ["قيود يدوية غير متوازنة", 3, 1901000, "أخطاء إدخال يدوية (مايو 2026)"],
    ["تعديل اشتراك على فني (معدّل)", 23, 264000, "بق مسار تعديل الاشتراك (يغيّر طرفاً واحداً)"],
    ["أجور مهام غير مُرحّلة", 3, 3000, "ضحايا نافذة تيتيم CompanyId"],
]
for i, row in enumerate(summ):
    for j, val in enumerate(row):
        c = ws.cell(4 + i, 1 + j, val)
        c.border = border
        if i == 0:
            c.fill = hdr_fill; c.font = hdr_font
        if j == 2 and i > 0:
            c.number_format = "#,##0"
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 8
ws.column_dimensions["C"].width = 15
ws.column_dimensions["D"].width = 55

# ---- أوراق التفاصيل ----
for name, fn, desc in sheets:
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws["A1"] = desc
    ws["A1"].font = title_font
    with open(os.path.join(base, fn), encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        continue
    header = rows[0]
    start = 3
    for j, col in enumerate(header, 1):
        c = ws.cell(start, j, col.replace("_", " "))
        c.fill = hdr_fill; c.font = hdr_font; c.border = border
        c.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows[1:], start + 1):
        for j, val in enumerate(row, 1):
            colname = header[j - 1]
            if colname in num_cols:
                try:
                    val = float(val)
                except (ValueError, TypeError):
                    pass
            c = ws.cell(r, j, val)
            c.border = border
            if colname in num_cols:
                c.number_format = "#,##0"
    # عرض الأعمدة
    for j, col in enumerate(header, 1):
        letter = openpyxl.utils.get_column_letter(j)
        maxlen = max([len(col)] + [len(str(rows[k][j - 1])) for k in range(1, len(rows))]) if len(rows) > 1 else len(col)
        ws.column_dimensions[letter].width = min(max(maxlen + 2, 12), 60)

wb.save(out)
print("SAVED:", out)
print("SHEETS:", [s.title for s in wb.worksheets])
